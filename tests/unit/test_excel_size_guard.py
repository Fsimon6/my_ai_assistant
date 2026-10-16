# -*- coding: utf-8 -*-
"""P-8 Excel 规模守卫回归：**守卫必须在任何全量 workbook DOM 之前**生效。

背景：守卫原先放在第一次 `load_workbook(read_only=False)` **之后**，而那份全量 DOM 正是
内存峰值的来源 —— 超限文件仍会先把内存打满，之后才（可能）报 413，即"守卫防不住 OOM"。

本文件锁定三件事：
  1. **限制语义不变**：仍用 EXCEL_MAX_SHEET_ROWS / EXCEL_MAX_SHEET_CELLS /
     EXCEL_MAX_TOTAL_CELLS，仍是 `>` 判定（limit 本身通过，limit+1 被拒）；
  2. **守卫先于全量加载**：超限文件在解析过程中**不会**出现 read_only=False 的加载
     （用 spy 包住 `openpyxl.load_workbook` 断言零调用）；
  3. **不改变既有行为**：正常文件仍走原有的两次全量加载，损坏文件仍由解析流程按原样报错
     （不会被预检改写成规模错误）。

测试用 `tests/helpers/fixtures.build_probe_xlsx()` 生成"极小体积 + 极大使用范围"的 xlsx，
因此无需生成真实的百万行数据即可精确测边界。
"""
from pathlib import Path

import openpyxl
import pytest

from backend.excel import parser as P
from tests.helpers import fixtures as F

LIMITS = P._excel_limits()


def _probe(tmp_path: Path, sheets):
    return F.build_probe_xlsx(tmp_path / 'probe.xlsx', sheets)


# ==========================================================================
# 1. 单 Sheet 行数上限
# ==========================================================================
def test_sheet_rows_over_limit_is_rejected(tmp_path):
    path = _probe(tmp_path, [('S1', LIMITS['rows'] + 1, 1)])     # 200_001 行 × 1 列
    with pytest.raises(P.ExcelSizeLimitError) as err:
        P._guard_xlsx_before_load(path)
    assert '超过单表上限' in str(err.value)


def test_sheet_rows_at_limit_passes(tmp_path):
    _probe(tmp_path, [('S1', LIMITS['rows'], 1)])               # 恰好等于上限 -> 通过
    P._guard_xlsx_before_load(tmp_path / 'probe.xlsx')


# ==========================================================================
# 2. 单 Sheet 单元格上限：limit-1 / limit / limit+1
# ==========================================================================
def test_sheet_cells_boundary(tmp_path):
    cols = 20                       # 2_000_000 / 20 = 100_000，可精确取到"恰好等于上限"
    exact = LIMITS['sheet_cells'] // cols

    P._guard_xlsx_before_load(_probe(tmp_path, [('S1', exact - 1, cols)]))   # limit-1
    P._guard_xlsx_before_load(_probe(tmp_path, [('S1', exact, cols)]))       # limit
    with pytest.raises(P.ExcelSizeLimitError) as err:
        P._guard_xlsx_before_load(_probe(tmp_path, [('S1', exact + 1, cols)]))   # limit+1
    assert '超过单表上限' in str(err.value)


# ==========================================================================
# 3. 整簿累计单元格上限
# ==========================================================================
def test_total_cells_over_limit_is_rejected(tmp_path):
    each = [('S1', 180_000, 10)]                    # 单 Sheet 1.8M < 2M，行 180k < 200k
    P._guard_xlsx_before_load(_probe(tmp_path, each))                # 单表通过
    three = [('S1', 180_000, 10), ('S2', 180_000, 10), ('S3', 180_000, 10)]
    with pytest.raises(P.ExcelSizeLimitError) as err:                # 5.4M > 5M
        P._guard_xlsx_before_load(_probe(tmp_path, three))
    assert '整个工作簿累计' in str(err.value)


# ==========================================================================
# 4. 守卫时机：全量 DOM 加载必须发生在守卫之后
# ==========================================================================
def test_over_limit_file_never_triggers_full_load(tmp_path, monkeypatch):
    """超限文件：解析过程必须在**任何** read_only=False 加载之前就失败。"""
    calls = []
    real_load = openpyxl.load_workbook

    def spy(*args, **kwargs):
        calls.append(kwargs.get('read_only'))
        return real_load(*args, **kwargs)

    monkeypatch.setattr(openpyxl, 'load_workbook', spy)
    path = _probe(tmp_path, [('S1', LIMITS['rows'] + 1, 1)])

    with pytest.raises(P.ExcelSizeLimitError):
        P.ExcelParser().parse(str(path), document_id='probe-size')

    assert calls == []          # 未发生任何 load_workbook（更没有 read_only=False 的全量加载）


def test_normal_file_still_uses_two_full_loads(tmp_path, monkeypatch):
    """正常文件：行为与改动前一致（仍是两次 read_only=False 全量加载）。"""
    calls = []
    real_load = openpyxl.load_workbook

    def spy(*args, **kwargs):
        calls.append(kwargs.get('read_only'))
        return real_load(*args, **kwargs)

    monkeypatch.setattr(openpyxl, 'load_workbook', spy)
    rep = F.parse_golden('golden_small.xlsx', document_id='golden-size')

    assert rep.sheet_count == 2
    assert calls == [False, False]
    assert rep.sheets[0].sheet_name == 'OrderSKUList'


def test_corrupt_file_keeps_existing_error_type(tmp_path):
    """损坏文件：预检静默跳过，不得把既有错误改写成规模错误。"""
    bad = tmp_path / 'broken.xlsx'
    bad.write_bytes(b'not a zip at all')

    P._guard_xlsx_before_load(bad)                  # 不抛规模错误

    with pytest.raises(Exception) as err:
        P.ExcelParser().parse(str(bad), document_id='broken')
    assert not isinstance(err.value, P.ExcelSizeLimitError)


# ==========================================================================
# 5. 缺 r 的 <row> / <c>：口径必须与 openpyxl DOM 完全一致
#    （复审发现的真实缺陷：<c> 无 r 时曾错误沿用上一行的行号）
# ==========================================================================
#: 刻意构造的异常形状（无 <dimension>，且 <row>/<c> 可能缺 r）
ODD_SHAPES = {
    'c 无 r（row 有 r）':
        '<row r="10"><c><v>1</v></c><c><v>2</v></c><c><v>3</v></c></row>',
    'row 无 r / c 有 r':
        '<row><c r="A1"><v>1</v></c></row><row><c r="B2"><v>2</v></c></row>',
    '混合 r 与无 r':
        '<row r="2"><c r="A2"><v>1</v></c><c><v>2</v></c></row>'
        '<row><c><v>3</v></c></row>',
    '同 row 连续无 r 的 c':
        '<row r="7"><c><v>1</v></c><c><v>2</v></c><c><v>3</v></c><c><v>4</v></c></row>',
    'row r + c r 跨行 + 无 r 的 c':
        '<row r="10"><c r="B20"><v>1</v></c><c><v>2</v></c></row>',
}
#: **独立**期望值（人工按 openpyxl 语义推算，不取自被测实现）
ODD_EXPECTED = {
    'c 无 r（row 有 r）': (10, 3),
    'row 无 r / c 有 r': (2, 2),
    '混合 r 与无 r': (3, 2),
    '同 row 连续无 r 的 c': (7, 4),
    'row r + c r 跨行 + 无 r 的 c': (20, 3),
}


def _scanner_dims(path):
    """直接用 ZIP/XML 扫描得到 [(sheet, rows, cols)]（与被测入口同源，但绕开 parse）。"""
    import zipfile

    with zipfile.ZipFile(str(path)) as zf:
        out = []
        for name, part in P._xlsx_worksheet_parts(zf):
            with zf.open(part) as stream:
                rows, cols = P._scan_sheet_used_range(stream, LIMITS, name)
            out.append((name, rows, cols))
        return out


def _dom_dims(path):
    """openpyxl DOM 的 (max_row, max_column)，作为**独立对照口径**。"""
    wb = openpyxl.load_workbook(filename=str(path), data_only=True, read_only=False)
    return [(ws.title, int(ws.max_row or 0), int(ws.max_column or 0)) for ws in wb.worksheets]


@pytest.mark.parametrize('label', sorted(ODD_SHAPES))
def test_scanner_matches_dom_for_odd_shapes(tmp_path, label):
    """缺 r 的 <row>/<c>：scanner 必须同时等于 DOM 与人工推算值。"""
    path = F.build_raw_xlsx(tmp_path / 'odd.xlsx', [('S1', ODD_SHAPES[label])])

    expect = [('S1',) + ODD_EXPECTED[label]]
    assert _scanner_dims(path) == expect          # 与人工推算值一致
    assert _dom_dims(path) == expect              # 且与 openpyxl DOM 实际行为一致


def test_odd_shape_over_limit_is_rejected_before_any_full_load(tmp_path, monkeypatch):
    """缺 r 的 <c>（行号只能靠 <row r> 推断）也必须被前置守卫拦下 —— 不允许发生全量加载。

    这是本轮缺陷的回归点：若行号推断错误（沿用上一行），该文件会被误判为 1 行而放行，
    随后进入 read_only=False 全量加载。
    """
    calls = []
    real_load = openpyxl.load_workbook

    def spy(*args, **kwargs):
        calls.append(kwargs.get('read_only'))
        return real_load(*args, **kwargs)

    monkeypatch.setattr(openpyxl, 'load_workbook', spy)
    body = '<row r="%d"><c><v>1</v></c></row>' % (LIMITS['rows'] + 1)
    path = F.build_raw_xlsx(tmp_path / 'odd_over.xlsx', [('S1', body)])

    with pytest.raises(P.ExcelSizeLimitError) as err:
        P.ExcelParser().parse(str(path), document_id='odd-over')

    assert '超过单表上限' in str(err.value)
    assert calls == []          # 护栏在解析之前就拒绝，未发生任何 load_workbook


def test_odd_shape_file_parses_normally(tmp_path):
    """正常规模的缺 r 形状：仍能被真实解析路径正常解析（前置预检不得误伤）。"""
    path = F.build_raw_xlsx(tmp_path / 'odd_ok.xlsx', [('S1', ODD_SHAPES['row 无 r / c 有 r'])])

    rep = P.ExcelParser().parse(str(path), document_id='odd-ok')

    assert rep.sheet_count == 1
    assert rep.sheets[0].column_count == 2          # A..B
    assert rep.sheets[0].row_count >= 1
